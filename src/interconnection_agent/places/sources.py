"""Readers for the operator's public files about places: bottlenecks, costs, upgrades.

  * Bottleneck list (2024): which connection points sit behind which overloaded lines or
    transformers, and how much room is left behind each.
  * Cost to add room (2026 estimates): for each bottleneck, the MW the next upgrade would add
    and its cost, in 2022 dollars.
  * Planned upgrades (July 2026 tracker, plus the yearly plan's fact sheets for cost ranges).

These readers only parse. Linking what they name to our places goes through the reviewed
tables in this package.
"""

from __future__ import annotations

import datetime
import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from pypdf import PdfReader

from interconnection_agent.ingest import _cells

BOTTLENECK_LIST = "attachment-b1-v8-constraint-mapping-2024-ipe.xlsx"
COST_TO_ADD_ROOM = (
    "attachment-a-transmission-capability-estimates-for-use-in-the-cpuc-irp-process-2026.xlsx"
)
UPGRADE_TRACKER = "approved-projects-transmission-planning-process-jul-2026.xlsx"
UPGRADE_FACT_SHEETS = "board-approved-2025-2026-transmission-plan-appendix-h-projects.pdf"


@dataclass(frozen=True)
class Upgrade:
    plan_id: str
    name: str
    utility: str | None
    expected_finish: datetime.date | None
    expected_finish_year: int | None
    cost_low_musd: float | None = None
    cost_high_musd: float | None = None
    fact_sheet_page: int | None = None


@dataclass(frozen=True)
class Sources:
    # Connection-point name (as the bottleneck list writes it) -> bottleneck names it's behind.
    behind: dict[str, list[str]]
    # Bottleneck name (as the bottleneck list writes it) -> MW of room left behind it.
    room_left_mw: dict[str, float]
    # Bottleneck name (as the cost file writes it) -> (MW the next upgrade adds, cost $M 2022).
    cost_to_add_room: dict[str, tuple[float, float]]
    upgrades: list[Upgrade]
    # Fact sheets whose cost couldn't be read, or whose name isn't in the upgrade tracker.
    unread_upgrade_costs: list[str]


def load_sources(data_dir: Path) -> Sources:
    behind, room = _bottleneck_list(data_dir / BOTTLENECK_LIST)
    upgrades, unread = _upgrades(data_dir / UPGRADE_TRACKER, data_dir / UPGRADE_FACT_SHEETS)
    return Sources(
        behind=behind,
        room_left_mw=room,
        cost_to_add_room=_cost_to_add_room(data_dir / COST_TO_ADD_ROOM),
        upgrades=upgrades,
        unread_upgrade_costs=unread,
    )


def _bottleneck_list(path: Path) -> tuple[dict[str, list[str]], dict[str, float]]:
    """Each sheet is a grid: connection points down the side, bottlenecks across the top,
    a tick where a point sits behind a bottleneck. Summary rows above the points give each
    bottleneck's room left ("Available TPD")."""
    behind: dict[str, list[str]] = {}
    room: dict[str, float] = {}
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            tops = [i for i, r in enumerate(rows) if str(r[0] or "").startswith("POI Name")]
            if not tops:
                raise ValueError(f"{path.name}, sheet {sheet.title}: no 'POI Name' header row")
            top = tops[0]
            names = {
                j: _cells.clean(v) for j, v in enumerate(rows[top]) if j >= 3 and _cells.clean(v)
            }
            for row in rows[top + 1 :]:
                label = _cells.clean(row[0])
                if label is None:
                    continue
                if label.startswith("Available TPD"):
                    for j, name in names.items():
                        mw = _cells.mw_or_none(row[j]) if j < len(row) else None
                        if name and mw is not None:
                            room[name] = mw
                    continue
                if row[1] is None:  # the other summary rows carry no area
                    continue
                ticks = [n for j, n in names.items() if n and j < len(row) and _cells.clean(row[j])]
                behind.setdefault(label, []).extend(ticks)
    finally:
        workbook.close()
    return behind, room


def _cost_to_add_room(path: Path) -> dict[str, tuple[float, float]]:
    """Rows: bottleneck, places affected, condition, room now (MW), MW the next upgrade adds,
    its description, its cost ($M, 2022 dollars), ... Section titles have no numbers."""
    costs: dict[str, tuple[float, float]] = {}
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["TxCapabilityEstimates_2026"]
        # Columns are read by position, so check the two header rows still say what we expect.
        top, sub = (list(r) for r in sheet.iter_rows(min_row=2, max_row=3, values_only=True))
        if not (
            str(top[0]).startswith("Transmission Constraint")
            and str(sub[4]).startswith("Incremental due to Network Upgrade")
            and str(sub[6]).startswith("Cost (2022$)")
        ):
            raise ValueError(f"{path.name}: column layout changed; update _cost_to_add_room")
        for row in sheet.iter_rows(min_row=4, values_only=True):
            name = _cells.clean(row[0])
            added = _cells.mw_or_none(row[4]) if len(row) > 6 else None
            cost = _cells.mw_or_none(row[6]) if len(row) > 6 else None
            if name and added and cost is not None:
                costs[name] = (added, cost)
    finally:
        workbook.close()
    return costs


def _upgrades(tracker: Path, fact_sheets: Path) -> tuple[list[Upgrade], list[str]]:
    costs, unread = _fact_sheet_costs(fact_sheets)
    upgrades: list[Upgrade] = []
    workbook = openpyxl.load_workbook(tracker, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets[1:]:  # the first sheet only explains the codes
            rows = sheet.iter_rows(values_only=True)
            header = [" ".join(str(h).split()) if h else "" for h in next(rows)]
            for row in rows:
                plan_id, name = _cells.clean(row[0]), _cells.clean(row[1])
                if not plan_id or not name:
                    continue
                finish, year = _latest_finish(header, row)
                low, high, page = costs.get(name, (None, None, None))
                upgrades.append(
                    Upgrade(plan_id, name, _cells.clean(row[2]), finish, year, low, high, page)
                )
    finally:
        workbook.close()
    tracked = {u.name for u in upgrades}
    unread += [f"{name}: not in the upgrade tracker" for name in costs if name not in tracked]
    return upgrades, unread


def _latest_finish(
    header: list[str], row: tuple[object, ...]
) -> tuple[datetime.date | None, int | None]:
    """The most recent expected finish: the "Current" column, else the latest earlier one,
    else the date at approval. Some cells give only a year."""
    order = [i for i, h in enumerate(header) if h.startswith("Current In-Service")]
    order += [
        i for i, h in reversed(list(enumerate(header))) if h.startswith("Previous In-Service")
    ]
    order += [
        i for i, h in enumerate(header) if h.lower().startswith("in-service date at approval")
    ]
    for i in order:
        value = row[i] if i < len(row) else None
        date = _cells.as_date(value)
        if date:
            return date, date.year
        text = _cells.clean(value)
        if text and re.fullmatch(r"\d{4}", text):
            return None, int(text)
    return None, None


def _fact_sheet_costs(path: Path) -> tuple[dict[str, tuple[float, float, int]], list[str]]:
    """Upgrade name -> (low, high) cost in $ millions, from the plan's one-page fact sheets.

    Costs are written many ways ("$31.3M - $62.6M", "$11-$15 million", "$44 M-89 M",
    "$1,270 - $1,800 million", "$114.8 million"). All the numbers are millions unless the
    text says billion. A sheet whose cost can't be read is returned in the second list.
    """
    costs: dict[str, tuple[float, float, int]] = {}
    unread: list[str] = []
    blocks = [
        (number, block)
        for number, page in enumerate(PdfReader(path).pages, start=1)
        for block in re.split(r"\n(?=Name )", page.extract_text() or "")
    ]
    for page_number, block in blocks:
        name_match = re.match(r"Name (.+?)\s*\n", block)
        if not name_match or "Project Cost" not in block:
            continue
        name = " ".join(name_match.group(1).split())
        cost_text = block.split("Project Cost", 1)[1].split("Alternatives", 1)[0]
        numbers = [float(n.replace(",", "")) for n in re.findall(r"\d[\d,]*(?:\.\d+)?", cost_text)]
        if not numbers or len(numbers) > 2:
            unread.append(f"{name}: cost '{' '.join(cost_text.split())}'")
            continue
        scale = 1000.0 if re.search(r"billion|\dB\b", cost_text, re.I) else 1.0
        costs[name] = (numbers[0] * scale, numbers[-1] * scale, page_number)
    return costs, unread
